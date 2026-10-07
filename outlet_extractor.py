"""
Outlet Extractor - simple desktop app
=====================================
Reads raw Excel report exports (title block on top, header somewhere in the
first rows, many columns) and converts them to a clean, fixed output structure.

Output structure
  - Default: the Outlet.xlsx layout
        TERRITORY | ROUTE | SEQUENCE | OUTLET CODE | OUTLET
  - Or upload ANY Excel file as the structure ("Upload structure file"): its
    header row defines the output columns (names, order, sheet name, column
    widths and header styling). Input columns are matched by header name.
    Columns not found in an input file are left blank (and reported).

Two output methods
  1) Separate files  - one clean file per input file
  2) One combined    - a single file containing all rows from all inputs

Before saving you can (all optional)
  - sort the rows (up to 3 levels, ascending / descending)
  - add NEW columns with your own titles (fixed value, blank, or row number)
  - set a report title (PDF heading, and optionally a title row in Excel)

Two output formats (tick one or both)
  - Excel (.xlsx)
  - PDF   (.pdf)   - formatted table, header repeated on every page

Requirements:  Python 3.9+  and  openpyxl, reportlab
               pip install openpyxl reportlab
Run:           python outlet_extractor.py
Build an .exe: pip install pyinstaller
               pyinstaller --onefile --windowed outlet_extractor.py
"""

import os
import re
import tkinter as tk
from copy import copy
from datetime import datetime
from itertools import chain, islice
from tkinter import filedialog, messagebox, ttk
from xml.sax.saxutils import escape

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

try:  # PDF support is optional so the Excel features still work without it
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

# --------------------------------------------------------------------------
# Settings
# --------------------------------------------------------------------------
KEY_COLUMN = "OUTLET CODE"      # if present, a blank value here marks a blank/footer row
HEADER_SEARCH_ROWS = 40         # how far down to look for the header row in input files
STRUCT_SEARCH_ROWS = 30         # ... and in an uploaded structure file
SORT_NONE = "(none)"
END = "(end)"                   # "insert new column at the end"
MAX_EXTRA = 5                   # max number of new columns


# --------------------------------------------------------------------------
# Output structure
# --------------------------------------------------------------------------
class Structure:
    """Describes the output table: column headers and how to draw them."""

    def __init__(self, headers, widths=None, sheet_name="Export", styles=None,
                 label="Default (Outlet.xlsx layout)", weights=None,
                 title="Outlet List", row_label="outlets", is_default=False):
        self.headers = list(headers)
        self.widths = widths or {}        # header -> Excel column width
        self.sheet_name = sheet_name
        self.styles = styles or {}        # header -> (font, fill, alignment, border)
        self.label = label
        self.weights = weights or {}      # header -> relative PDF width
        self.title = title                # default PDF heading
        self.row_label = row_label
        self.is_default = is_default


_DEFAULT_HEADERS = ["TERRITORY", "ROUTE", "SEQUENCE", "OUTLET CODE", "OUTLET"]
DEFAULT_STRUCTURE = Structure(
    _DEFAULT_HEADERS,
    widths=dict(zip(_DEFAULT_HEADERS, [15.3, 29.4, 11.7, 17.6, 23.4])),
    weights=dict(zip(_DEFAULT_HEADERS, [12, 26, 13, 19, 30])),
    is_default=True,
)
TARGET_COLUMNS = DEFAULT_STRUCTURE.headers   # kept for convenience


def _norm(value):
    """Normalise a header cell for matching."""
    return re.sub(r"\s+", " ", str(value).strip().upper()) if value is not None else ""


def load_structure(path):
    """Build a Structure from the header row of an Excel file (data rows are ignored)."""
    wb = load_workbook(path)
    try:
        ws = wb.active
        best_row, best_score = None, 0
        for r in range(1, min(ws.max_row, STRUCT_SEARCH_ROWS) + 1):
            score = sum(1 for c in ws[r] if isinstance(c.value, str) and c.value.strip())
            if score > best_score:          # first row with the most text cells
                best_row, best_score = r, score
        if best_row is None:
            raise ValueError("Could not find a header row in that file.")

        headers, widths, styles, seen = [], {}, {}, set()
        for c in ws[best_row]:
            if c.value is None or str(c.value).strip() == "":
                continue
            h = re.sub(r"\s+", " ", str(c.value).strip())
            if h.upper() in seen:
                raise ValueError(f"The structure file has the column '{h}' more than once.")
            seen.add(h.upper())
            headers.append(h)
            ltr = get_column_letter(c.column)
            dim = ws.column_dimensions[ltr] if ltr in ws.column_dimensions else None
            if dim is not None and dim.width:
                widths[h] = float(dim.width)
            styles[h] = (copy(c.font), copy(c.fill), copy(c.alignment), copy(c.border))

        stem = os.path.splitext(os.path.basename(path))[0]
        return Structure(
            headers, widths=widths, sheet_name=ws.title, styles=styles,
            label=f"{os.path.basename(path)}  (sheet '{ws.title}', header on row {best_row})",
            weights={h: w for h, w in widths.items()},
            title=f"{stem} List", row_label="rows",
        )
    finally:
        wb.close()


# --------------------------------------------------------------------------
# Core logic (no GUI) - easy to test on its own
# --------------------------------------------------------------------------
def extract(path, structure=None):
    """Read one input file. Returns (rows, missing_headers).
    Each row follows structure.headers; columns not found in the file are blank."""
    structure = structure or DEFAULT_STRUCTURE
    headers = structure.headers
    wanted = {_norm(h) for h in headers}

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        head = list(islice(it, HEADER_SEARCH_ROWS))

        best_i, best_n, best_normed = None, 0, None
        for i, row in enumerate(head):
            normed = [_norm(v) for v in row]
            n = len(wanted.intersection(normed))
            if n > best_n:                       # first row with the most matches
                best_i, best_n, best_normed = i, n, normed
        if best_i is None:
            raise ValueError("No header row matching the output structure was found "
                             f"(looked for: {', '.join(headers[:6])}{'…' if len(headers) > 6 else ''}).")

        col_map = {h: (best_normed.index(_norm(h)) if _norm(h) in best_normed else None)
                   for h in headers}
        missing = [h for h, i in col_map.items() if i is None]
        key_i = next((i for h, i in col_map.items() if _norm(h) == KEY_COLUMN and i is not None), None)

        out = []
        for row in chain(head[best_i + 1:], it):
            rec = []
            for h in headers:
                i = col_map[h]
                v = row[i] if (i is not None and i < len(row)) else None
                if isinstance(v, str):
                    v = v.strip()
                if isinstance(v, float) and v.is_integer():
                    v = int(v)  # 5.0 -> 5
                rec.append(v)
            if key_i is not None:
                key = row[key_i] if key_i < len(row) else None
                if key is None or str(key).strip() == "":
                    continue  # blank / footer row
            elif all(_is_blank(v) for v in rec):
                continue
            out.append(rec)
        return out, missing
    finally:
        wb.close()


def extract_rows(path, structure=None):
    """Rows only (see extract())."""
    return extract(path, structure)[0]


# ---- sorting ---------------------------------------------------------------
def _is_blank(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


def _sort_key(v):
    """Numbers sort numerically; text sorts case-insensitively with natural
    number order (so 'ST 2' comes before 'ST 10')."""
    if isinstance(v, (int, float)):
        return (0, v, [])
    parts = re.split(r"(\d+)", str(v).strip().casefold())
    return (1, 0, [int(p) if p.isdigit() else p for p in parts])


def sort_rows(rows, levels, headers=None):
    """levels = [(column_name, descending_bool), ...] - first item is the main sort.
    Blank cells always go to the bottom. Returns a new list."""
    headers = headers or DEFAULT_STRUCTURE.headers
    rows = list(rows)
    for col, desc in reversed(levels):          # stable sorts, last key first
        i = headers.index(col)
        filled = [r for r in rows if not _is_blank(r[i])]
        blank = [r for r in rows if _is_blank(r[i])]
        filled.sort(key=lambda r: _sort_key(r[i]), reverse=desc)
        rows = filled + blank
    return rows


def describe_sort(levels):
    return ", ".join(f"{c} {'descending' if d else 'ascending'}" for c, d in levels)


# ---- extra (user-defined) columns -----------------------------------------
def _fill_value(value, n):
    """'' -> blank cell, '{n}' -> row number, otherwise the text (with {n} replaced)."""
    value = (value or "").strip()
    if value == "":
        return None
    if value == "{n}":
        return n
    return value.replace("{n}", str(n))


def validate_extras(extras, headers=None):
    """Return an error message, or None if the new column titles are fine."""
    used = {c.casefold() for c in (headers or DEFAULT_STRUCTURE.headers)}
    for ex in extras or []:
        t = (ex.get("title") or "").strip()
        if not t:
            continue
        if t.casefold() in used:
            return f"The column title '{t}' is used more than once (or already exists)."
        used.add(t.casefold())
    return None


def apply_extras(rows, extras, headers=None):
    """Insert the user's new columns. extras = [{'title','value','after'}, ...]
    Returns (headers, new_rows). Rows with an empty title are ignored."""
    base = list(headers or DEFAULT_STRUCTURE.headers)
    err = validate_extras(extras, base)
    if err:
        raise ValueError(err)
    headers = list(base)
    out = [list(r) for r in rows]
    for ex in extras or []:
        title = (ex.get("title") or "").strip()
        if not title:
            continue
        after = ex.get("after", END)
        pos = headers.index(after) + 1 if after in headers else len(headers)
        headers.insert(pos, title)
        for n, r in enumerate(out, start=1):
            r.insert(pos, _fill_value(ex.get("value"), n))
    return headers, out


# ---- writers ---------------------------------------------------------------
def write_output(rows, out_path, headers=None, title=None, structure=None):
    """Write rows to an .xlsx following the structure (sheet name, widths, header style).
    If `title` is given it goes in row 1 and the table header moves to row 3."""
    structure = structure or DEFAULT_STRUCTURE
    headers = headers or structure.headers
    wb = Workbook()
    ws = wb.active
    ws.title = structure.sheet_name
    header_row = 1
    if title:
        ws.append([title])
        ws["A1"].font = Font(name="Calibri", size=14, bold=True)
        ws.append([])
        header_row = 3
    ws.append(headers)
    for cell, h in zip(ws[header_row], headers):
        style = structure.styles.get(h)
        if style:
            cell.font, cell.fill, cell.alignment, cell.border = (copy(s) for s in style)
        else:
            cell.font = Font(name="Calibri", size=11, bold=True)
    for rec in rows:
        ws.append(rec)
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = (
            structure.widths.get(h) or max(15, len(h) + 4))
    wb.save(out_path)


def _is_numeric_col(rows, i):
    vals = [r[i] for r in rows if not _is_blank(r[i])]
    return bool(vals) and all(isinstance(v, (int, float)) for v in vals)


def write_pdf(rows, out_path, title="Outlet List", headers=None, structure=None):
    """Write rows to a formatted PDF table (header repeats on each page).
    Switches to landscape automatically when there are many columns."""
    if not PDF_AVAILABLE:
        raise RuntimeError("PDF output needs reportlab:  pip install reportlab")

    structure = structure or DEFAULT_STRUCTURE
    headers = headers or structure.headers
    page = landscape(A4) if len(headers) > 7 else A4
    left = right = 14 * mm
    doc = SimpleDocTemplate(
        out_path, pagesize=page, leftMargin=left, rightMargin=right,
        topMargin=16 * mm, bottomMargin=16 * mm, title=title,
    )
    usable = page[0] - left - right

    title_st = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=15, leading=19)
    sub_st = ParagraphStyle("s", fontName="Helvetica", fontSize=9, leading=12,
                            textColor=colors.HexColor("#555555"))
    head_st = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=8.5, leading=10,
                             textColor=colors.white)
    head_c = ParagraphStyle("hc", parent=head_st, alignment=1)
    cell_st = ParagraphStyle("c", fontName="Helvetica", fontSize=8.5, leading=10.5)
    num_st = ParagraphStyle("n", parent=cell_st, alignment=1)  # centred

    sub = f"{len(rows)} {structure.row_label}"
    route_i = next((i for i, h in enumerate(headers) if _norm(h) == "ROUTE"), None)
    if route_i is not None:
        n_routes = len({r[route_i] for r in rows})
        sub += f"  |  {n_routes} route{'s' if n_routes != 1 else ''}"
    sub += f"  |  Generated {datetime.now():%Y-%m-%d %H:%M}"

    def cell(v, style):
        return Paragraph(escape("" if v is None else str(v)), style)

    centred = [_is_numeric_col(rows, i) for i in range(len(headers))]
    data = [[Paragraph(escape(h), head_c if c else head_st) for h, c in zip(headers, centred)]]
    for r in rows:
        data.append([cell(v, num_st if c else cell_st) for v, c in zip(r, centred)])

    weights = [structure.weights.get(h) or 10 for h in headers]
    total = sum(weights)
    widths = [usable * w / total for w in weights]
    # never let a column be narrower than its header text (avoids "TERRITOR/Y")
    mins = [stringWidth(h, "Helvetica-Bold", 8.5) + 14 for h in headers]
    for _ in range(5):
        widths = [max(w, m) for w, m in zip(widths, mins)]
        over = sum(widths) - usable
        if over <= 0.5:
            break
        flex = [w - m for w, m in zip(widths, mins)]
        tf = sum(flex)
        if tf <= 0:
            break
        widths = [w - over * f / tf for w, f in zip(widths, flex)]
    table = Table(data, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5F9")]),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#D0D7E2")),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#777777"))
        canvas.drawRightString(page[0] - right, 9 * mm, f"Page {d.page}")
        canvas.drawString(left, 9 * mm, title)
        canvas.restoreState()

    story = [Paragraph(escape(title), title_st), Paragraph(escape(sub), sub_st),
             Spacer(1, 6 * mm), table]
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def natural_key(path):
    """Sort 'Outlet (2).xlsx' before 'Outlet (10).xlsx'."""
    name = os.path.basename(path).lower()
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", name)]


def _save(rows, base_path, title, make_xlsx, make_pdf, log, extras, xl_title, structure):
    """Add the new columns, then write as .xlsx and/or .pdf. Returns (written, failed)."""
    headers, full_rows = apply_extras(rows, extras, structure.headers)
    ok = bad = 0
    jobs = []
    if make_xlsx:
        jobs.append((".xlsx", lambda p: write_output(full_rows, p, headers, xl_title, structure)))
    if make_pdf:
        jobs.append((".pdf", lambda p: write_pdf(full_rows, p, title, headers, structure)))
    for ext, fn in jobs:
        path = base_path + ext
        try:
            fn(path)
            ok += 1
            log(f"          -> {os.path.basename(path)}")
        except PermissionError:
            bad += 1
            log(f"  FAILED  cannot write {os.path.basename(path)} (is it open in another program?)")
        except Exception as e:
            bad += 1
            log(f"  FAILED  {os.path.basename(path)}: {e}")
    return ok, bad


def process(files, mode, out_dir, combined_name, log, make_xlsx=True, make_pdf=False,
            sort_levels=None, extra_cols=None, doc_title="", title_in_excel=False,
            structure=None):
    """mode = 'separate' | 'combined'. Returns number of output files written."""
    structure = structure or DEFAULT_STRUCTURE
    headers = structure.headers
    files = sorted(files, key=natural_key)
    all_rows, written, failed = [], 0, 0
    sort_levels = sort_levels or []
    extra_cols = [e for e in (extra_cols or []) if (e.get("title") or "").strip()]
    doc_title = (doc_title or "").strip()

    err = validate_extras(extra_cols, headers)
    if err:
        raise ValueError(err)
    if not structure.is_default:
        log(f"Output structure: {structure.label}")
        log("Columns: " + " | ".join(headers))
    if sort_levels:
        log(f"Sorting by: {describe_sort(sort_levels)}")
    if extra_cols:
        log("New columns: " + ", ".join(e["title"].strip() for e in extra_cols))
    if doc_title:
        log(f"Report title: {doc_title}")
    if (not structure.is_default) or sort_levels or extra_cols or doc_title:
        log("")

    route_i = next((i for i, h in enumerate(headers) if _norm(h) == "ROUTE"), None)
    key_i = next((i for i, h in enumerate(headers) if _norm(h) == KEY_COLUMN), None)

    for f in files:
        name = os.path.basename(f)
        try:
            rows, missing = extract(f, structure)
        except Exception as e:  # keep going with other files
            failed += 1
            log(f"  FAILED  {name}: {e}")
            continue

        log(f"  OK      {name}  ({len(rows)} rows)")
        if missing:
            log(f"  WARNING not found in this file, left blank: {', '.join(missing)}")
        if mode == "separate":
            stem = os.path.splitext(name)[0]
            routes = sorted({r[route_i] for r in rows if r[route_i]}) if route_i is not None else []
            route_part = routes[0] if len(routes) == 1 else None
            base_title = doc_title or structure.title
            title = f"{base_title} - {route_part}" if route_part else base_title
            if sort_levels:
                rows = sort_rows(rows, sort_levels, headers)
            ok, bad = _save(rows, os.path.join(out_dir, f"{stem}_structured"), title,
                            make_xlsx, make_pdf, log, extra_cols,
                            title if (doc_title and title_in_excel) else None, structure)
            written += ok
            failed += bad
        else:
            all_rows.extend(rows)

    if mode == "combined" and all_rows:
        if sort_levels:
            all_rows = sort_rows(all_rows, sort_levels, headers)
        base = re.sub(r"\.(xlsx|pdf)$", "", combined_name, flags=re.I)
        title = doc_title or (f"{structure.title} - All Routes" if route_i is not None
                              else structure.title)
        log(f"\n  Combined: {len(all_rows)} rows")
        ok, bad = _save(all_rows, os.path.join(out_dir, base), title,
                        make_xlsx, make_pdf, log, extra_cols,
                        title if (doc_title and title_in_excel) else None, structure)
        written += ok
        failed += bad

        # duplicate check on the key column (default: OUTLET CODE)
        if key_i is not None:
            codes = [r[key_i] for r in all_rows]
            dups = {c for c in codes if c is not None and codes.count(c) > 1}
            if dups:
                log(f"  NOTE    {len(dups)} duplicate {headers[key_i]} value(s) found, "
                    f"e.g. {sorted(dups, key=str)[:3]}")

    log(f"\nDone. {written} file(s) written, {failed} problem(s).")
    return written


# --------------------------------------------------------------------------
# GUI
# --------------------------------------------------------------------------
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Outlet Extractor")
        self.geometry("800x780")
        self.minsize(720, 720)

        self.structure = DEFAULT_STRUCTURE
        self.files = []
        self.mode = tk.StringVar(value="combined")
        self.out_dir = tk.StringVar()
        self.combined_name = tk.StringVar(value="Outlet_combined")
        self.sort_vars = [(tk.StringVar(value=SORT_NONE), tk.StringVar(value="Ascending"))
                          for _ in range(3)]
        self.sort_boxes = []
        self.report_title = tk.StringVar()
        self.title_in_excel = tk.BooleanVar(value=False)
        self.extra_vars = []        # [{'title','value','after'} -> StringVars]
        self.want_xlsx = tk.BooleanVar(value=True)
        self.want_pdf = tk.BooleanVar(value=PDF_AVAILABLE)

        pad = {"padx": 10, "pady": 5}

        # ---------------- tabs ----------------
        self.nb = ttk.Notebook(self, height=430)
        self.nb.pack(fill="x", padx=10, pady=(10, 4))
        self.tab_files = ttk.Frame(self.nb)
        self.tab_custom = ttk.Frame(self.nb)
        self.tab_out = ttk.Frame(self.nb)
        self.nb.add(self.tab_files, text="  1. Input files & structure  ")
        self.nb.add(self.tab_custom, text="  2. Sort & extra columns  ")
        self.nb.add(self.tab_out, text="  3. Output  ")
        self._build_files_tab(pad)
        self._build_custom_tab(pad)
        self._build_output_tab(pad)

        # ---------------- run + log ----------------
        self.run_btn = ttk.Button(self, text="Extract", command=self.run)
        self.run_btn.pack(pady=6)

        f4 = ttk.LabelFrame(self, text="Log")
        f4.pack(fill="both", expand=True, **pad)
        self.log_box = tk.Text(f4, height=6, state="disabled", wrap="none")
        lsb = ttk.Scrollbar(f4, command=self.log_box.yview)
        self.log_box.config(yscrollcommand=lsb.set)
        self.log_box.pack(side="left", fill="both", expand=True, padx=(6, 0), pady=6)
        lsb.pack(side="right", fill="y", pady=6)

        self._toggle()
        self._apply_structure(DEFAULT_STRUCTURE)

    # ---- tab builders ----
    def _build_files_tab(self, pad):
        t = self.tab_files
        btns = ttk.Frame(t)
        btns.pack(fill="x", padx=8, pady=8)
        ttk.Button(btns, text="Add files…", command=self.add_files).pack(side="left")
        ttk.Button(btns, text="Add folder…", command=self.add_folder).pack(side="left", padx=6)
        ttk.Button(btns, text="Remove selected", command=self.remove_selected).pack(side="left")
        ttk.Button(btns, text="Clear", command=self.clear_files).pack(side="left", padx=6)
        self.count_lbl = ttk.Label(btns, text="0 files")
        self.count_lbl.pack(side="right")

        # structure box is packed first at the bottom so the list can expand above it
        fst = ttk.LabelFrame(t, text="Output structure")
        fst.pack(side="bottom", fill="x", padx=8, pady=(0, 8))
        self.struct_lbl = ttk.Label(fst, text="", justify="left", wraplength=720)
        self.struct_lbl.pack(anchor="w", padx=8, pady=(6, 4))
        sb_row = ttk.Frame(fst)
        sb_row.pack(fill="x", padx=8, pady=(0, 8))
        ttk.Button(sb_row, text="Upload structure file…", command=self.pick_structure).pack(side="left")
        ttk.Button(sb_row, text="Reset to default", command=self.reset_structure).pack(side="left", padx=6)
        ttk.Button(sb_row, text="Check input match", command=self.check_mapping).pack(side="left")

        lb_frame = ttk.Frame(t)
        lb_frame.pack(fill="both", expand=True, padx=8, pady=(0, 8))
        self.listbox = tk.Listbox(lb_frame, selectmode="extended")
        sb = ttk.Scrollbar(lb_frame, command=self.listbox.yview)
        self.listbox.config(yscrollcommand=sb.set)
        self.listbox.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")

    def _build_custom_tab(self, pad):
        t = self.tab_custom

        # sorting
        fs = ttk.LabelFrame(t, text="Sort rows before saving (optional)")
        fs.pack(fill="x", padx=8, pady=(8, 4))
        for n, (col_var, dir_var) in enumerate(self.sort_vars):
            ttk.Label(fs, text="Sort by" if n == 0 else "Then by", width=8).grid(
                row=n, column=0, padx=(8, 2), pady=3, sticky="w")
            cb = ttk.Combobox(fs, textvariable=col_var, values=[SORT_NONE],
                              state="readonly", width=18)
            cb.grid(row=n, column=1, padx=4, pady=3)
            self.sort_boxes.append(cb)
            ttk.Combobox(fs, textvariable=dir_var, values=["Ascending", "Descending"],
                         state="readonly", width=12).grid(row=n, column=2, padx=4, pady=3)
        ttk.Button(fs, text="Reset sorting", command=self.reset_sort).grid(
            row=0, column=3, padx=20)

        # title + new columns
        fc = ttk.LabelFrame(t, text="Report title & new columns (optional)")
        fc.pack(fill="both", expand=True, padx=8, pady=4)

        rt = ttk.Frame(fc)
        rt.pack(fill="x", padx=8, pady=(6, 0))
        ttk.Label(rt, text="Report title:").pack(side="left")
        ttk.Entry(rt, textvariable=self.report_title, width=44).pack(side="left", padx=6)
        ttk.Checkbutton(fc, text="Also show the title in the Excel file (table then starts on row 3)",
                        variable=self.title_in_excel).pack(anchor="w", padx=8, pady=(2, 0))

        ttk.Label(
            fc, foreground="#555", justify="left",
            text="New columns: the value is filled into every row.\n"
                 "Leave it empty for a blank column, or type {n} for a row number (1, 2, 3 ...).",
        ).pack(anchor="w", padx=8, pady=(6, 2))

        self.extra_box = ttk.Frame(fc)
        self.extra_box.pack(fill="x", padx=8, pady=2)
        self.add_col_btn = ttk.Button(fc, text="+ Add column", command=self.add_extra)
        self.add_col_btn.pack(anchor="w", padx=8, pady=(2, 8))

    def _build_output_tab(self, pad):
        t = self.tab_out
        f2 = ttk.LabelFrame(t, text="Output method")
        f2.pack(fill="x", padx=8, pady=(8, 4))
        ttk.Radiobutton(
            f2, text="Method 1 - Separate files (one structured file per input)",
            variable=self.mode, value="separate", command=self._toggle,
        ).pack(anchor="w", padx=8, pady=(6, 0))
        ttk.Radiobutton(
            f2, text="Method 2 - One combined file (all inputs merged into a single sheet)",
            variable=self.mode, value="combined", command=self._toggle,
        ).pack(anchor="w", padx=8, pady=(0, 4))
        row = ttk.Frame(f2)
        row.pack(fill="x", padx=8, pady=(0, 8))
        ttk.Label(row, text="Combined file name:").pack(side="left")
        self.name_entry = ttk.Entry(row, textvariable=self.combined_name, width=32)
        self.name_entry.pack(side="left", padx=6)

        ff = ttk.LabelFrame(t, text="Output format (tick one or both)")
        ff.pack(fill="x", padx=8, pady=4)
        rf = ttk.Frame(ff)
        rf.pack(fill="x", padx=8, pady=6)
        ttk.Checkbutton(rf, text="Excel (.xlsx)", variable=self.want_xlsx).pack(side="left")
        pdf_cb = ttk.Checkbutton(rf, text="PDF (.pdf)", variable=self.want_pdf)
        pdf_cb.pack(side="left", padx=20)
        if not PDF_AVAILABLE:
            pdf_cb.config(state="disabled")
            ttk.Label(rf, text="(install reportlab to enable:  pip install reportlab)",
                      foreground="#a33").pack(side="left")

        f3 = ttk.LabelFrame(t, text="Output folder")
        f3.pack(fill="x", padx=8, pady=4)
        r3 = ttk.Frame(f3)
        r3.pack(fill="x", padx=6, pady=8)
        ttk.Entry(r3, textvariable=self.out_dir).pack(side="left", fill="x", expand=True)
        ttk.Button(r3, text="Browse…", command=self.pick_out).pack(side="left", padx=(6, 0))

    # ---- helpers ----
    def _toggle(self):
        state = "normal" if self.mode.get() == "combined" else "disabled"
        self.name_entry.config(state=state)

    def log(self, msg):
        self.log_box.config(state="normal")
        self.log_box.insert("end", msg + "\n")
        self.log_box.see("end")
        self.log_box.config(state="disabled")
        self.update_idletasks()

    def clear_log(self):
        self.log_box.config(state="normal")
        self.log_box.delete("1.0", "end")
        self.log_box.config(state="disabled")

    def _refresh_list(self):
        self.listbox.delete(0, "end")
        for f in self.files:
            self.listbox.insert("end", f)
        n = len(self.files)
        self.count_lbl.config(text=f"{n} file{'s' if n != 1 else ''}")

    def _add(self, paths):
        for p in paths:
            base = os.path.basename(p)
            if base.startswith("~$") or not base.lower().endswith((".xlsx", ".xlsm")):
                continue
            if p not in self.files:
                self.files.append(p)
        self.files.sort(key=natural_key)
        self._refresh_list()

    # ---- output structure ----
    def _apply_structure(self, st):
        """Switch to a new output structure and refresh everything that depends on it."""
        self.structure = st
        shown = " | ".join(st.headers[:12]) + (" | …" if len(st.headers) > 12 else "")
        if st.is_default:
            self.struct_lbl.config(text=f"Using the default structure (Outlet.xlsx layout):\n{shown}")
        else:
            self.struct_lbl.config(
                text=f"Using structure from: {st.label}\n{len(st.headers)} columns: {shown}")
        choices = [SORT_NONE] + st.headers
        for cb in self.sort_boxes:
            cb.config(values=choices)
        for col_var, _ in self.sort_vars:
            if col_var.get() not in choices:
                col_var.set(SORT_NONE)
        for ev in self.extra_vars:
            if ev["after"].get() not in [END] + st.headers:
                ev["after"].set(END)
        self._render_extras()

    def load_structure_file(self, path):
        """Load a structure from a file path (used by the button and by tests)."""
        st = load_structure(path)
        self._apply_structure(st)
        return st

    def pick_structure(self):
        path = filedialog.askopenfilename(
            title="Select an Excel file that has the output structure (header row)",
            filetypes=[("Excel files", "*.xlsx *.xlsm")],
        )
        if not path:
            return
        try:
            self.load_structure_file(path)
        except Exception as e:
            messagebox.showerror("Structure file", f"Could not read the structure:\n{e}")

    def reset_structure(self):
        self._apply_structure(DEFAULT_STRUCTURE)

    def check_mapping(self):
        """Report how well each input file matches the current structure (nothing is saved)."""
        if not self.files:
            messagebox.showwarning("No files", "Please add at least one input file first.")
            return
        self.clear_log()
        self.log(f"Checking {len(self.files)} file(s) against the structure…\n")
        for f in self.files:
            name = os.path.basename(f)
            try:
                rows, missing = extract(f, self.structure)
            except Exception as e:
                self.log(f"  FAILED   {name}: {e}")
                continue
            if missing:
                self.log(f"  PARTIAL  {name}: {len(rows)} rows; blank (not found): {', '.join(missing)}")
            else:
                self.log(f"  MATCH    {name}: {len(rows)} rows, all columns found")
        self.log("\nCheck finished (nothing was saved).")

    # ---- new-columns editor ----
    def add_extra(self):
        if len(self.extra_vars) >= MAX_EXTRA:
            return
        self.extra_vars.append({"title": tk.StringVar(), "value": tk.StringVar(),
                                "after": tk.StringVar(value=END)})
        self._render_extras()

    def remove_extra(self, idx):
        del self.extra_vars[idx]
        self._render_extras()

    def _render_extras(self):
        for w in self.extra_box.winfo_children():
            w.destroy()
        if self.extra_vars:
            for c, txt in enumerate(["Column title", "Value", "Insert after"]):
                ttk.Label(self.extra_box, text=txt, foreground="#555").grid(
                    row=0, column=c, sticky="w", padx=2)
        for i, ev in enumerate(self.extra_vars, start=1):
            ttk.Entry(self.extra_box, textvariable=ev["title"], width=20).grid(
                row=i, column=0, padx=2, pady=2)
            ttk.Entry(self.extra_box, textvariable=ev["value"], width=24).grid(
                row=i, column=1, padx=2, pady=2)
            ttk.Combobox(self.extra_box, textvariable=ev["after"],
                         values=[END] + self.structure.headers,
                         state="readonly", width=14).grid(row=i, column=2, padx=2, pady=2)
            ttk.Button(self.extra_box, text="✕", width=3,
                       command=lambda k=i - 1: self.remove_extra(k)).grid(row=i, column=3, padx=2)
        self.add_col_btn.config(state="normal" if len(self.extra_vars) < MAX_EXTRA else "disabled")

    def get_extras(self):
        return [{"title": ev["title"].get().strip(), "value": ev["value"].get(),
                 "after": ev["after"].get()}
                for ev in self.extra_vars if ev["title"].get().strip()]

    # ---- actions ----
    def add_files(self):
        paths = filedialog.askopenfilenames(
            title="Select input Excel files",
            filetypes=[("Excel files", "*.xlsx *.xlsm")],
        )
        self._add(paths)
        if paths and not self.out_dir.get():
            self.out_dir.set(os.path.dirname(paths[0]))

    def add_folder(self):
        d = filedialog.askdirectory(title="Select a folder with Excel files")
        if d:
            self._add(os.path.join(d, n) for n in os.listdir(d))
            if not self.out_dir.get():
                self.out_dir.set(d)

    def reset_sort(self):
        for col_var, dir_var in self.sort_vars:
            col_var.set(SORT_NONE)
            dir_var.set("Ascending")

    def get_sort_levels(self):
        levels, seen = [], set()
        for col_var, dir_var in self.sort_vars:
            col = col_var.get()
            if col != SORT_NONE and col not in seen:   # ignore repeated columns
                seen.add(col)
                levels.append((col, dir_var.get() == "Descending"))
        return levels

    def remove_selected(self):
        for i in reversed(self.listbox.curselection()):
            del self.files[i]
        self._refresh_list()

    def clear_files(self):
        self.files.clear()
        self._refresh_list()

    def pick_out(self):
        d = filedialog.askdirectory(title="Select output folder")
        if d:
            self.out_dir.set(d)

    def run(self):
        if not self.files:
            messagebox.showwarning("No files", "Please add at least one input file.")
            return
        out_dir = self.out_dir.get().strip()
        if not out_dir:
            messagebox.showwarning("No output folder", "Please choose an output folder (Output tab).")
            return
        if self.mode.get() == "combined" and not self.combined_name.get().strip():
            messagebox.showwarning("File name", "Please enter a name for the combined file.")
            return
        if not (self.want_xlsx.get() or self.want_pdf.get()):
            messagebox.showwarning("Output format", "Tick Excel and/or PDF (Output tab).")
            return
        extras = self.get_extras()
        err = validate_extras(extras, self.structure.headers)
        if err:
            messagebox.showwarning("New columns", err)
            return
        os.makedirs(out_dir, exist_ok=True)

        self.run_btn.config(state="disabled")
        self.clear_log()
        self.log(f"Processing {len(self.files)} file(s)…\n")
        try:
            n = process(self.files, self.mode.get(), out_dir,
                        self.combined_name.get().strip(), self.log,
                        make_xlsx=self.want_xlsx.get(),
                        make_pdf=self.want_pdf.get() and PDF_AVAILABLE,
                        sort_levels=self.get_sort_levels(),
                        extra_cols=extras,
                        doc_title=self.report_title.get(),
                        title_in_excel=self.title_in_excel.get(),
                        structure=self.structure)
            if n:
                if messagebox.askyesno("Finished", "Done! Open the output folder?"):
                    self._open_folder(out_dir)
        except Exception as e:
            messagebox.showerror("Error", str(e))
        finally:
            self.run_btn.config(state="normal")

    @staticmethod
    def _open_folder(path):
        try:
            if os.name == "nt":
                os.startfile(path)
            elif os.uname().sysname == "Darwin":
                os.system(f'open "{path}"')
            else:
                os.system(f'xdg-open "{path}"')
        except Exception:
            pass


if __name__ == "__main__":
    App().mainloop()
